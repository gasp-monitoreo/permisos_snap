"""Django Ninja API — todos los endpoints de la plataforma SNAP."""
import io
import json as _json
import logging
import os
import re as _re
import unicodedata
import time
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import List, Optional

import requests as _requests
from django.conf import settings
from django.http import FileResponse, HttpResponse
from ninja import File, Form, NinjaAPI, UploadedFile

from .core import (
    INVESTIGACION_PROCESO_IDS, api_get, es_tramite_visible,
    fetch_all_processes, fetch_all_tramites, normalize_tramite,
)
from .email_utils import (
    GMAIL_AVAILABLE, _GMAIL_IMPORT_ERROR, email_html_body,
    get_gmail_creds, gmail_creds_file, create_gmail_draft, send_via_smtp,
)
from .excel_utils import build_compliance_excel, build_tramites_excel
from .fetcher import get_compliance_data, get_portada_historico_data, refresh_caches

INV_DIR       = settings.INVESTIGACIONES_DIR
INV_EXTENSIONS = {'.pdf', '.doc', '.docx', '.csv', '.xls', '.xlsx'}
DOMAIN        = 'https://conaf.cerofilas.gob.cl'

from .usuarios import SesionAuth, perfil_de, requiere
from . import privacidad

log = logging.getLogger(__name__)

# Toda la API exige sesión iniciada (antes era pública y exponía datos personales)
api = NinjaAPI(title='SNAP API', version='2.0', docs_url=None, auth=SesionAuth())


def _region_forzada(request):
    """Región a la que el perfil restringe los datos (None = todo el país)."""
    return perfil_de(request.user).region_forzada()


def _filtrar_region(request, lista):
    reg = _region_forzada(request)
    return lista if not reg else [t for t in lista if reg in (t.get('regiones_list') or [])]


# ── Helpers ───────────────────────────────────────────────────────────────────

def _inv_archivos(base_path):
    result = []
    for root, dirs, files in os.walk(base_path):
        dirs.sort()
        for f in sorted(files):
            if f.startswith('.'):
                continue
            if os.path.splitext(f)[1].lower() not in INV_EXTENSIONS:
                continue
            full = os.path.join(root, f)
            rel  = os.path.relpath(full, INV_DIR).replace('\\', '/')
            result.append({
                'nombre': f,
                'url': '/investigaciones/' + rel,
                'size_kb': round(os.path.getsize(full) / 1024, 1),
            })
    return result


def _norm_nombre_carpeta(nombre: str) -> str:
    s = unicodedata.normalize('NFD', nombre or '')
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    s = _re.sub(r'[^\w\s]', '', s)
    return s.strip().replace(' ', '_') or 'Sin_Nombre'


def _carpeta_existente(nombre_carpeta: str):
    if not os.path.isdir(INV_DIR):
        return None
    tokens = nombre_carpeta.lower().split('_')
    for d in os.listdir(INV_DIR):
        if all(tok in d.lower() for tok in tokens if tok):
            return d
    return None


# ── Procesos ──────────────────────────────────────────────────────────────────

@api.get('/procesos')
def get_procesos(request):
    try:
        return {'ok': True, 'data': fetch_all_processes()}
    except Exception as e:
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=500)


# ── Trámites ──────────────────────────────────────────────────────────────────

@api.get('/tramites')
def get_tramites(
    request,
    proceso_id: Optional[int] = None,
    created_at_start: Optional[str] = None,
    created_at_end: Optional[str]   = None,
    updated_at_start: Optional[str] = None,
    updated_at_end: Optional[str]   = None,
    ended_at_start: Optional[str]   = None,
    ended_at_end: Optional[str]     = None,
    fresco: bool = False,
):
    try:
        # Rango de años del panel → caché (ver tramites_cache.py); fresco=true fuerza descarga
        solo_rango = created_at_start and created_at_end and not any(
            [proceso_id, updated_at_start, updated_at_end, ended_at_start, ended_at_end])
        if solo_rango:
            from . import tramites_cache
            desde, hasta = tramites_cache.rango_anios(created_at_start, created_at_end)
            forzar = fresco and perfil_de(request.user).puede('editar')
            entrada = tramites_cache.obtener(desde, hasta, forzar=forzar)
            data = privacidad.tramites(request, _filtrar_region(request, entrada['data']))
            return {'ok': True, 'data': data, 'total': len(data),
                    'actualizado_en': datetime.fromtimestamp(entrada['ts']).isoformat()}

        date_cutoff = None
        if created_at_start:
            try:
                date_cutoff = datetime.fromtimestamp(int(created_at_start))
            except Exception:
                pass
        raw = fetch_all_tramites(
            proceso_id=proceso_id,
            created_at_start=created_at_start, created_at_end=created_at_end,
            updated_at_start=updated_at_start, updated_at_end=updated_at_end,
            ended_at_start=ended_at_start, ended_at_end=ended_at_end,
            date_cutoff=date_cutoff,
        )
        normalized = [normalize_tramite(t) for t in raw]
        normalized = _filtrar_region(request, [t for t in normalized if es_tramite_visible(t)])
        normalized = privacidad.tramites(request, normalized)
        return {'ok': True, 'data': normalized, 'total': len(normalized)}
    except Exception as e:
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=500)


@api.get('/tramites/{tramite_id}')
def get_tramite(request, tramite_id: int):
    try:
        data = api_get(f'/tramites/{tramite_id}')
        t    = data.get('tramite', data)
        norm = normalize_tramite(t)
        if not _filtrar_region(request, [norm]):
            return api.create_response(request, {'ok': False, 'error': 'Trámite fuera de su región.'}, status=403)
        return {'ok': True, 'data': privacidad.tramites(request, [norm])[0]}
    except _requests.HTTPError as e:
        status = e.response.status_code if e.response is not None else 500
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=status)
    except Exception as e:
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=500)


# ── Cache ─────────────────────────────────────────────────────────────────────

@api.post('/cache/refresh')
@requiere(accion='editar')
def cache_refresh(request):
    refresh_caches()
    return {'ok': True, 'msg': 'Regeneración en background iniciada.'}


# ── Portada historico ─────────────────────────────────────────────────────────

@api.get('/portada/historico')
def portada_historico(request):
    try:
        data = privacidad.tramites(request, _filtrar_region(request, get_portada_historico_data()))
        return {'ok': True, 'data': data, 'total': len(data)}
    except Exception as e:
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=500)


# ── Compliance ────────────────────────────────────────────────────────────────

@api.get('/compliance/investigacion')
def compliance_investigacion(request):
    try:
        resultados, revisados = get_compliance_data()
        resultados = privacidad.tramites(request, _filtrar_region(request, resultados))
        return {'ok': True, 'data': resultados, 'total': len(resultados),
                'candidatos_revisados': revisados}
    except Exception as e:
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=500)


# ── Export Excel ──────────────────────────────────────────────────────────────

@api.get('/export/excel')
@requiere(seccion='exportar')
def export_excel(
    request,
    proceso_id: Optional[int]   = None,
    created_at_start: Optional[str] = None,
    created_at_end: Optional[str]   = None,
    ended_at_start: Optional[str]   = None,
    ended_at_end: Optional[str]     = None,
    estado: Optional[str]           = None,
):
    try:
        raw = fetch_all_tramites(
            proceso_id=proceso_id,
            created_at_start=created_at_start, created_at_end=created_at_end,
            ended_at_start=ended_at_start, ended_at_end=ended_at_end,
        )
        tramites = [normalize_tramite(t) for t in raw]
        tramites = privacidad.tramites(request, _filtrar_region(request, [t for t in tramites if es_tramite_visible(t)]))
        if estado:
            tramites = [t for t in tramites if t['estado'] == estado]

        try:
            procesos = fetch_all_processes()
            proceso_map = {
                (p.get('proceso', p).get('id') if 'proceso' in p else p.get('id')):
                (p.get('proceso', p).get('nombre') if 'proceso' in p else p.get('nombre'))
                for p in procesos
            }
        except Exception:
            proceso_map = {}

        buf      = build_tramites_excel(tramites, proceso_map)
        filename = f"tramites_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        resp = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        return resp
    except Exception as e:
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=500)


@api.get('/export/compliance/excel')
@requiere(seccion='exportar')
def export_compliance_excel(request):
    try:
        resultados, _ = get_compliance_data()
        buf      = build_compliance_excel(privacidad.tramites(request, _filtrar_region(request, resultados)))
        filename = f"incumplimiento_numeral6_{datetime.now().strftime('%Y%m%d')}.xlsx"
        resp = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        return resp
    except Exception as e:
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=500)


# ── Adjuntos ──────────────────────────────────────────────────────────────────

def _candidatos_archivo(archivo, tramite_id=''):
    TOKEN = settings.CEROFILAS_TOKEN
    API   = settings.CEROFILAS_API_BASE
    return [c for c in [
        f'{DOMAIN}/uploads/datos/{archivo}',            # donde Cero Filas publica los adjuntos
        f'{DOMAIN}/backend/storage/{archivo}',
        f'{DOMAIN}/backend/storage/app/public/{archivo}',
        f'{DOMAIN}/storage/{archivo}',
        f'{DOMAIN}/storage/app/public/{archivo}',
        f'{API}/archivos/{archivo}',
        f'{API}/uploads/{archivo}',
        (f'{API}/tramites/{tramite_id}/archivos/{archivo}' if tramite_id else None),
        f'{API}/files/{archivo}',
        f'{API}/download/{archivo}',
    ] if c]


def _url_de_cerofilas(url):
    """True solo si la URL es https y su host es exactamente el de CeroFilas.
    El token de la API se envía a estas URLs: nunca a otro host."""
    from urllib.parse import urlparse
    try:
        u, base = urlparse(url), urlparse(DOMAIN)
    except ValueError:
        return False
    return u.scheme == 'https' and (u.hostname or '').lower() == base.hostname and u.port in (None, 443)         and not u.username and not u.password


def _descargar_url(url):
    if not _url_de_cerofilas(url):
        return None, None
    TOKEN = settings.CEROFILAS_TOKEN
    for kwargs in [
        {'params': {'token': TOKEN}},
        {'headers': {'Authorization': f'Bearer {TOKEN}'}},
    ]:
        try:
            r = _requests.get(url, timeout=30, **kwargs)
            if r.status_code == 200:
                ct = r.headers.get('Content-Type', 'application/octet-stream')
                if 'text/html' not in ct:
                    return r, ct
        except Exception:
            pass
    return None, None


@api.get('/adjunto')
@requiere(superusuario=True)
def descargar_adjunto(
    request,
    url: Optional[str]      = None,
    archivo: Optional[str]  = None,
    tramite_id: Optional[str] = None,
    nombre: str             = 'archivo',
):
    if not url and not archivo:
        return api.create_response(request, {'ok': False, 'error': 'Se requiere url o archivo'}, status=400)
    if url:
        if not _url_de_cerofilas(url):
            return api.create_response(request, {'ok': False, 'error': 'URL no permitida'}, status=403)
        candidatos = [url]
    else:
        candidatos = _candidatos_archivo(archivo, tramite_id or '')

    for candidate in candidatos:
        resp, ct = _descargar_url(candidate)
        if resp is not None:
            http_resp = HttpResponse(resp.content, content_type=ct)
            http_resp['Content-Disposition'] = f'attachment; filename="{nombre}"'
            return http_resp

    return api.create_response(request, {'ok': False, 'error': f"No se pudo descargar '{archivo}'"}, status=502)


@api.get('/adjunto/debug')
@requiere(superusuario=True)
def debug_adjunto(request, archivo: str, tramite_id: Optional[str] = None):
    TOKEN = settings.CEROFILAS_TOKEN
    if not archivo:
        return api.create_response(request, {'ok': False, 'error': 'Falta archivo'}, status=400)
    resultados = []
    for url in _candidatos_archivo(archivo, tramite_id or ''):
        for kwargs, modo in [
            ({'params': {'token': TOKEN}}, 'token_param'),
            ({'headers': {'Authorization': f'Bearer {TOKEN}'}}, 'bearer_header'),
        ]:
            try:
                r = _requests.get(url, timeout=10, **kwargs)
                resultados.append({'url': url, 'modo': modo, 'status': r.status_code,
                                   'content_type': r.headers.get('Content-Type', '')})
            except Exception as e:
                resultados.append({'url': url, 'modo': modo, 'error': str(e)})
    return {'ok': True, 'archivo': archivo, 'resultados': resultados}


# ── Search by area ────────────────────────────────────────────────────────────

@api.get('/search/area')
@requiere(seccion='busqueda')
def search_by_area(
    request, q: str,
    year_from: Optional[int] = None,
    year_to: Optional[int]   = None,
):
    q = q.strip().lower()
    if len(q) < 2:
        return api.create_response(request, {'ok': False, 'error': 'Ingrese al menos 2 caracteres'}, status=400)

    now = datetime.now()
    year_from = max(2019, min(year_from or now.year, now.year))
    year_to   = max(year_from, min(year_to or now.year, now.year))

    created_at_start = str(int(datetime(year_from, 1, 1).timestamp()))
    created_at_end   = str(int(datetime(year_to, 12, 31, 23, 59, 59).timestamp()))

    raw_list = fetch_all_tramites(created_at_start=created_at_start, created_at_end=created_at_end)
    if not raw_list:
        raw_list = fetch_all_tramites()
    tramite_ids = [t.get('id') for t in raw_list if t.get('id')]

    def fetch_and_check(tid):
        try:
            data   = api_get(f'/tramites/{tid}')
            t      = data.get('tramite', data)
            norm   = normalize_tramite(t)
            datos  = norm.get('datos_raw', {})
            areas  = ' '.join(norm.get('areas_protegidas', [])).lower()
            for k, v in datos.items():
                if k.lower().startswith('region_de_'):
                    areas += ' ' + str(v).lower()
            return norm if q in areas else None
        except Exception:
            return None

    results = []
    if tramite_ids:
        with ThreadPoolExecutor(max_workers=min(30, len(tramite_ids))) as ex:
            futs = {ex.submit(fetch_and_check, tid): tid for tid in tramite_ids}
            for fut in as_completed(futs):
                r = fut.result()
                if r:
                    results.append(r)

    results = privacidad.tramites(request, _filtrar_region(request, [t for t in results if es_tramite_visible(t)]))
    results.sort(key=lambda x: x.get('fecha_inicio') or '', reverse=True)
    return {'ok': True, 'data': results, 'total': len(results),
            'searched': len(tramite_ids), 'year_from': year_from, 'year_to': year_to}


# ── Email ─────────────────────────────────────────────────────────────────────

@api.post('/email/send')
@requiere(accion='agregar', seccion='exportar')
def send_email_route(request):
    if not settings.ENVIO_CORREOS_ACTIVO:
        return api.create_response(request, {'ok': False, 'error': 'El envío de correos está desactivado.'}, status=403)
    import smtplib
    data = request.body and __import__('json').loads(request.body) or {}
    smtp = data.get('smtp', {})
    em   = data.get('email', {})
    missing = [k for k in ('smtp_server','smtp_port','smtp_user','smtp_pass','from_email') if not smtp.get(k)]
    if missing:
        return api.create_response(request, {'ok': False, 'error': f"Faltan campos SMTP: {', '.join(missing)}"}, status=400)
    if not em.get('to_email'):
        return api.create_response(request, {'ok': False, 'error': 'Falta email del destinatario'}, status=400)
    try:
        send_via_smtp(smtp, em)
        return {'ok': True}
    except smtplib.SMTPAuthenticationError:
        return api.create_response(request, {'ok': False, 'error': 'Contraseña incorrecta. Gmail requiere Contraseña de Aplicación.'}, status=401)
    except smtplib.SMTPConnectError as e:
        return api.create_response(request, {'ok': False, 'error': f'No se pudo conectar al servidor SMTP: {e}'}, status=503)
    except Exception as e:
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=500)


@api.get('/auth/gmail/status')
def gmail_auth_status(request):
    if not GMAIL_AVAILABLE:
        err = f' ({_GMAIL_IMPORT_ERROR})' if _GMAIL_IMPORT_ERROR else ''
        return {'ok': False, 'error': f'Librerías Google no disponibles{err}'}
    if not os.path.exists(gmail_creds_file()):
        return {'ok': False, 'error': 'Falta gmail_auth/credentials.json en el servidor'}
    creds = get_gmail_creds()
    return {'ok': True, 'authenticated': creds is not None}


@api.post('/email/draft')
@requiere(accion='agregar', seccion='exportar')
def create_draft_route(request):
    if not settings.ENVIO_CORREOS_ACTIVO:
        return api.create_response(request, {'ok': False, 'error': 'El envío de correos está desactivado.'}, status=403)
    data = request.body and __import__('json').loads(request.body) or {}
    em   = data.get('email', {})
    if not GMAIL_AVAILABLE:
        return api.create_response(request, {'ok': False, 'error': 'Librerías Google no instaladas'}, status=500)
    if not os.path.exists(gmail_creds_file()):
        return api.create_response(request, {'ok': False, 'error': 'NO_CREDS'}, status=400)
    if not em.get('to_email'):
        return api.create_response(request, {'ok': False, 'error': 'Falta email del destinatario'}, status=400)
    try:
        draft_id, err = create_gmail_draft(em)
        if err:
            return api.create_response(request, {'ok': False, 'error': err}, status=401)
        return {'ok': True, 'draft_id': draft_id}
    except Exception as e:
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=500)


# ── Investigaciones ───────────────────────────────────────────────────────────

@api.get('/investigaciones/drive/')
def investigaciones_drive(request):
    """Lista archivos desde Google Drive (carpeta SNAP configurada)."""
    from .drive_utils import get_cached_index
    data, err = get_cached_index()
    if err == 'needs_auth':
        return {
            'ok': False, 'needs_auth': True, 'data': {},
            'error': 'Drive no autorizado. Ve a /auth/gmail/start para autenticar.',
        }
    if err:
        return api.create_response(
            request, {'ok': False, 'needs_auth': False, 'data': {}, 'error': err}, status=500,
        )
    return {'ok': True, 'needs_auth': False, 'data': data}


@api.post('/investigaciones/drive/refresh')
@requiere(accion='editar')
def investigaciones_drive_refresh(request):
    """Invalida el caché de Drive y fuerza recarga."""
    from .drive_utils import invalidate_cache, get_cached_index
    invalidate_cache()
    data, err = get_cached_index()
    if err:
        return {'ok': False, 'error': err}
    return {'ok': True, 'data': data}


# ── Presión sobre áreas protegidas ────────────────────────────────────────────

@api.get('/presion')
def presion_areas(request, region: str, tipo: Optional[str] = None):
    """Permisos activos (en ejecución, programados, en evaluación) por área protegida de una región.
    tipo: 'inv' (investigación) o 'film' (filmación y fotografía); sin tipo, ambos."""
    from .presion import presion_region
    forzada = _region_forzada(request)
    if forzada and region != forzada:
        return api.create_response(request, {'ok': False, 'error': 'Solo puede consultar su región.'}, status=403)
    if tipo not in (None, 'inv', 'film'):
        return api.create_response(request, {'ok': False, 'error': 'tipo debe ser inv o film'}, status=400)
    try:
        return {'ok': True, 'region': region, 'tipo': tipo, 'data': _proteger_areas(request, presion_region(region, tipo))}
    except Exception as e:
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=500)


# ── Fotos de portada ──────────────────────────────────────────────────────────

_fotos_job = {'corriendo': False, 'ultimo': None}


def _es_superusuario(request):
    return bool(getattr(request, 'user', None) and request.user.is_authenticated and request.user.is_superuser)


_SOLO_SUPERUSUARIO = {'ok': False, 'error': 'Solo un superusuario puede gestionar las fotos.'}


@api.get('/fotos')
def fotos_lista(request, todas: bool = False):
    from .fotos import load_manifest
    fotos = load_manifest()['fotos']
    if not (todas and _es_superusuario(request)):   # las ocultas solo las ve el superusuario
        fotos = [f for f in fotos if f.get('activa', True)]
    return {'ok': True, 'data': [{**f, 'url': f'/fotos/{f["archivo"]}'} for f in fotos],
            'job': _fotos_job}


@api.post('/fotos/{archivo}/activa')
def fotos_toggle(request, archivo: str, activa: bool):
    if not _es_superusuario(request):
        return api.create_response(request, _SOLO_SUPERUSUARIO, status=403)
    from .fotos import _lock, load_manifest, save_manifest
    with _lock:
        m = load_manifest()
        for f in m['fotos']:
            if f['archivo'] == archivo:
                f['activa'] = activa
        save_manifest(m)
    return {'ok': True}


@api.post('/fotos/extraer')
def fotos_extraer(request):
    """Lanza en background la extracción de fotos desde los documentos Drive."""
    if not _es_superusuario(request):
        return api.create_response(request, _SOLO_SUPERUSUARIO, status=403)
    import threading
    from .drive_utils import get_cached_index, get_drive_service
    from . import fotos

    if _fotos_job['corriendo']:
        return {'ok': True, 'job': _fotos_job}
    svc, err = get_drive_service()
    if err:
        return {'ok': False, 'error': err}

    def run():
        _fotos_job.update({'corriendo': True, 'ultimo': None})
        try:
            index, e = get_cached_index()
            n = fotos.extraer(svc, index or {}, log=lambda *_: None) if not e else 0
            _fotos_job['ultimo'] = e or f'{n} fotos nuevas'
        except Exception as exc:
            _fotos_job['ultimo'] = f'error: {exc}'
        finally:
            _fotos_job['corriendo'] = False

    threading.Thread(target=run, daemon=True).start()
    return {'ok': True, 'job': {'corriendo': True}}


@api.get('/investigaciones/index')
def investigaciones_index(request):
    if not os.path.isdir(INV_DIR):
        return {'ok': True, 'data': {}}
    result = {}
    for investigador in sorted(os.listdir(INV_DIR)):
        inv_path = os.path.join(INV_DIR, investigador)
        if not os.path.isdir(inv_path) or investigador.startswith(('.', '_')):
            continue
        sub_dirs = [d for d in os.listdir(inv_path)
                    if os.path.isdir(os.path.join(inv_path, d)) and d.isdigit()]
        if sub_dirs:
            ids_data = {tid: _inv_archivos(os.path.join(inv_path, tid))
                        for tid in sorted(sub_dirs, key=int)}
            result[investigador] = {'tipo': 'nested', 'ids': ids_data}
        else:
            result[investigador] = {'tipo': 'flat', 'archivos': _inv_archivos(inv_path)}
    return {'ok': True, 'data': result}


@api.get('/investigaciones/archivos/{nombre}/{tramite_id}')
def investigaciones_archivos_id(request, nombre: str, tramite_id: str):
    carpeta = os.path.join(INV_DIR, nombre, tramite_id)
    if not os.path.abspath(carpeta).startswith(os.path.abspath(INV_DIR)):
        return api.create_response(request, {'ok': False}, status=403)
    if not os.path.isdir(carpeta):
        return {'ok': True, 'archivos': []}
    return {'ok': True, 'archivos': _inv_archivos(carpeta)}


@api.post('/investigaciones/upload')
@requiere(accion='agregar')
def investigaciones_upload(
    request,
    nombre: Form[str],
    tramite_id: Form[str],
    archivos: List[UploadedFile] = File(None),
):
    if not nombre or not tramite_id:
        return api.create_response(request, {'ok': False, 'error': 'Faltan parámetros'}, status=400)
    tramite_id = str(tramite_id).strip()
    if not tramite_id.isdigit() or len(tramite_id) > 12:
        return api.create_response(request, {'ok': False, 'error': 'ID de trámite inválido'}, status=400)
    if len(archivos or []) > 20:
        return api.create_response(request, {'ok': False, 'error': 'Máximo 20 archivos por envío'}, status=400)

    nombre_carpeta = _norm_nombre_carpeta(nombre)
    existente      = _carpeta_existente(nombre_carpeta)
    carpeta_base   = existente if existente else nombre_carpeta
    destino        = os.path.join(INV_DIR, carpeta_base, str(tramite_id))
    os.makedirs(destino, exist_ok=True)

    guardados, errores = [], []
    for f in (archivos or []):
        if not f or not f.name:
            continue
        fname = _re.sub(r'[^\w\.\-\(\) ]', '_', f.name).strip()
        if not fname:
            continue
        ext = os.path.splitext(fname)[1].lower()
        if ext not in INV_EXTENSIONS:
            errores.append(f'{f.name}: extensión no permitida ({ext})')
            continue
        if f.size > 50 * 1024 * 1024:
            errores.append(f'{f.name}: supera 50 MB')
            continue
        try:
            dest = os.path.join(destino, fname)
            if not os.path.abspath(dest).startswith(os.path.abspath(INV_DIR) + os.sep):
                errores.append(f'{fname}: ruta no permitida')
                continue
            with open(dest, 'wb') as out:
                for chunk in f.chunks():
                    out.write(chunk)
            guardados.append(fname)
        except Exception as exc:
            errores.append(f'{fname}: {exc}')

    return {'ok': True, 'guardados': guardados, 'errores': errores,
            'carpeta': f'{carpeta_base}/{tramite_id}'}


# ── Traducción ────────────────────────────────────────────────────────────────

@api.post('/translate')
def translate_text(request):
    import hashlib
    from django.core.cache import cache
    try:
        data = _json.loads(request.body or b'{}')
    except ValueError:
        data = {}
    texto = str(data.get('text') or '').strip()[:500]
    if not texto:
        return {'ok': True, 'traduccion': texto}
    clave = 'trad:' + hashlib.sha1(texto.encode()).hexdigest()
    if (guardada := cache.get(clave)) is not None:
        return {'ok': True, 'traduccion': guardada}
    if cache.get('trad:caido'):             # tras una falla, no insistir por 10 minutos
        return {'ok': False, 'error': 'Traductor no disponible', 'traduccion': texto}
    params = {'q': texto, 'langpair': 'en|es'}
    if settings.MYMEMORY_EMAIL:             # cuota diaria mayor en MyMemory
        params['de'] = settings.MYMEMORY_EMAIL
    try:
        data = _requests.get('https://api.mymemory.translated.net/get', params=params, timeout=6).json()
        if int(data.get('responseStatus') or 0) != 200:     # p. ej. cuota agotada: no mostrar el aviso como traducción
            cache.set('trad:caido', 1, 600)
            return {'ok': False, 'error': 'Traductor no disponible', 'traduccion': texto}
        traduccion = (data.get('responseData') or {}).get('translatedText') or texto
        cache.set(clave, traduccion, 30 * 24 * 3600)
        return {'ok': True, 'traduccion': traduccion}
    except Exception:
        cache.set('trad:caido', 1, 600)
        log.warning('Falla del traductor', exc_info=True)
        return {'ok': False, 'error': 'Traductor no disponible', 'traduccion': texto}


# ── Usuarios (solo superusuario) ─────────────────────────────────────────────

@api.get('/me')
def mi_perfil(request):
    u = request.user
    return {'ok': True, 'nombre': u.get_full_name() or u.username, 'email': u.email or u.username,
            **perfil_de(u).como_dict()}


@api.post('/me/tutorial')
def mi_tutorial(request):
    """El usuario elige no volver a ver el aviso del tutorial (o volver a verlo)."""
    p = perfil_de(request.user)
    p.ver_tutorial = not bool(_cuerpo(request).get('ocultar'))
    p.save(update_fields=['ver_tutorial', 'actualizado_en'])
    return {'ok': True, 'ver_tutorial': p.ver_tutorial}


def _usuario_dict(u):
    p = perfil_de(u)
    return {
        'id': u.id, 'email': u.email or u.username, 'nombre': u.get_full_name(),
        'activo': u.is_active, 'superusuario': u.is_superuser,
        'tipo': p.tipo, 'region': p.region, 'solo_su_region': p.solo_su_region,
        'secciones': p.secciones_visibles(),
        'permisos': {a: p.puede(a) for a in ('agregar', 'editar', 'borrar')},
        'clave_pendiente': p.debe_cambiar_password,
        'ultimo_ingreso': u.last_login.isoformat() if u.last_login else None,
    }


def _cuerpo(request):
    import json
    try:
        return json.loads(request.body or b'{}')
    except ValueError:
        return {}


def _aplicar_datos(u, p, d):
    """Copia los campos editables del formulario al usuario y su perfil."""
    from .models import SECCIONES
    from .usuarios import REGIONES
    if 'nombre' in d:
        partes = str(d['nombre']).split()
        u.first_name, u.last_name = ' '.join(partes[:1])[:150], ' '.join(partes[1:])[:150]
    if 'activo' in d and not u.is_superuser:
        u.is_active = bool(d['activo'])
    if 'tipo' in d and d['tipo'] in ('inv', 'film', 'ambos', 'nacional'):
        p.tipo = d['tipo']
    if 'region' in d:
        p.region = d['region'] if d['region'] in REGIONES else ''
    if 'solo_su_region' in d:
        p.solo_su_region = bool(d['solo_su_region']) and bool(p.region)
    if 'secciones' in d:
        p.secciones = [s for s in d['secciones'] if s in SECCIONES]
    for a in ('agregar', 'editar', 'borrar'):
        if a in (d.get('permisos') or {}):
            setattr(p, f'puede_{a}', bool(d['permisos'][a]))


@api.get('/usuarios')
@requiere(superusuario=True)
def usuarios_lista(request):
    from django.contrib.auth import get_user_model
    from .models import SECCIONES
    from .usuarios import REGIONES
    us = get_user_model().objects.order_by('-is_superuser', 'perfil__region', 'last_name', 'first_name')
    return {'ok': True, 'data': [_usuario_dict(u) for u in us],
            'secciones': SECCIONES, 'regiones': REGIONES}


@api.post('/usuarios')
@requiere(superusuario=True)
def usuarios_crear(request):
    from django.contrib.auth import get_user_model
    from .models import PerfilUsuario, secciones_por_defecto
    from .usuarios import restablecer_clave
    d = _cuerpo(request)
    email = str(d.get('email') or '').strip().lower()
    if '@' not in email:
        return api.create_response(request, {'ok': False, 'error': 'Correo inválido.'}, status=400)
    User = get_user_model()
    if User.objects.filter(username__iexact=email).exists():
        return api.create_response(request, {'ok': False, 'error': 'Ya existe un usuario con ese correo.'}, status=400)
    u = User.objects.create_user(username=email, email=email)
    tipo = d.get('tipo') if d.get('tipo') in ('inv', 'film', 'ambos', 'nacional') else 'nacional'
    p = PerfilUsuario.objects.create(user=u, tipo=tipo, secciones=secciones_por_defecto(tipo))
    _aplicar_datos(u, p, d)
    u.save()
    p.save()
    clave = restablecer_clave(u)
    return {'ok': True, 'usuario': _usuario_dict(u), 'clave_temporal': clave}


@api.put('/usuarios/{int:uid}')
@requiere(superusuario=True)
def usuarios_editar(request, uid: int):
    from django.contrib.auth import get_user_model
    u = get_user_model().objects.filter(id=uid).first()
    if not u:
        return api.create_response(request, {'ok': False, 'error': 'Usuario no encontrado.'}, status=404)
    p = perfil_de(u)
    _aplicar_datos(u, p, _cuerpo(request))
    u.save()
    p.save()
    return {'ok': True, 'usuario': _usuario_dict(u)}


@api.delete('/usuarios/{int:uid}')
@requiere(superusuario=True)
def usuarios_borrar(request, uid: int):
    from django.contrib.auth import get_user_model
    u = get_user_model().objects.filter(id=uid).first()
    if not u:
        return api.create_response(request, {'ok': False, 'error': 'Usuario no encontrado.'}, status=404)
    if u.id == request.user.id or u.is_superuser:
        return api.create_response(request, {'ok': False, 'error': 'No se puede eliminar un superusuario.'}, status=400)
    u.delete()
    return {'ok': True}


@api.post('/usuarios/{int:uid}/restablecer')
@requiere(superusuario=True)
def usuarios_restablecer(request, uid: int):
    """Clave temporal nueva (el usuario la olvidó). Se muestra una vez; la planilla la regenera."""
    from django.contrib.auth import get_user_model
    from .usuarios import restablecer_clave
    u = get_user_model().objects.filter(id=uid).first()
    if not u or u.id == request.user.id:
        return api.create_response(request, {'ok': False, 'error': 'Usuario no válido.'}, status=400)
    return {'ok': True, 'clave_temporal': restablecer_clave(u), 'usuario': _usuario_dict(u)}


@api.post('/usuarios/encargados')
@requiere(superusuario=True)
def usuarios_encargados(request):
    """Detecta los encargados regionales en los trámites y crea las cuentas que falten."""
    from .usuarios import crear_cuentas_encargados, detectar_encargados
    enc = detectar_encargados()
    creados, actualizados = crear_cuentas_encargados(enc)
    return {'ok': True, 'encargados': enc, 'creados': creados, 'actualizados': actualizados}


@api.post('/usuarios/planilla')
@requiere(superusuario=True)
def usuarios_planilla(request):
    """Planilla de usuarios con claves temporales, cifrada con la clave del superusuario.

    Cada descarga regenera la clave de quienes aún no definen la suya; quienes ya la
    cambiaron aparecen como «clave propia» y no se modifican.
    """
    import io
    import time as _t
    from datetime import datetime as _dt
    from concurrent.futures import ThreadPoolExecutor as _Pool
    from django.contrib.auth import get_user_model
    from django.contrib.auth.hashers import make_password
    from django.http import HttpResponse
    from msoffcrypto.format.ooxml import OOXMLFile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from .models import SECCIONES, TIPOS_PERFIL
    from .usuarios import clave_temporal, intento_bloqueado, registrar_intento

    yo = request.user
    if intento_bloqueado(yo.id):
        return api.create_response(request, {'ok': False, 'error': 'Demasiados intentos. Espere 10 minutos.'}, status=429)
    clave_admin = str(_cuerpo(request).get('password') or '')
    if not yo.check_password(clave_admin):
        registrar_intento(yo.id, False)
        _t.sleep(0.8)
        return api.create_response(request, {'ok': False, 'error': 'Contraseña incorrecta.'}, status=401)
    registrar_intento(yo.id, True)

    tipos = dict(TIPOS_PERFIL)
    usuarios = list(get_user_model().objects.select_related('perfil').order_by('perfil__region', 'last_name', 'first_name'))
    # Claves nuevas para quienes aún no definen la suya (se regeneran en cada descarga).
    # El hash (PBKDF2) toma ~1 s por clave: se calcula en paralelo.
    nuevas = {u.id: clave_temporal() for u in usuarios
              if u.id != yo.id and u.is_active and perfil_de(u).debe_cambiar_password}
    with _Pool(max_workers=8) as pool:
        hashes = dict(zip(nuevas, pool.map(make_password, nuevas.values())))
    for u in usuarios:
        if u.id in hashes:
            u.password = hashes[u.id]
            u.save(update_fields=['password'])

    filas = []
    for u in usuarios:
        p = perfil_de(u)
        clave = nuevas.get(u.id, '')
        estado = ('Inactivo' if not u.is_active else
                  'Clave temporal pendiente' if p.debe_cambiar_password else 'Activa – clave propia')
        filas.append([
            'Superusuario' if u.is_superuser else tipos.get(p.tipo, p.tipo), p.region or 'Nacional',
            u.get_full_name(), u.email or u.username, estado, clave,
            ', '.join(SECCIONES[s] for s in p.secciones_visibles()),
            'Sí', *('Sí' if p.puede(a) else 'No' for a in ('agregar', 'editar', 'borrar')),
            'Sí' if p.region_forzada() else 'No',
            u.last_login.strftime('%d-%m-%Y %H:%M') if u.last_login else 'Nunca',
        ])

    wb = Workbook()
    ws = wb.active
    ws.title = 'Léeme'
    for fila in [
        ['Planilla de usuarios · Permisos SNAP'], [],
        ['Generada', _dt.now().strftime('%d-%m-%Y %H:%M')], ['Por', yo.email or yo.username],
        ['Ingreso', f'https://{request.get_host()}/login/'], [],
        ['• Las claves temporales cambian cada vez que se descarga esta planilla: use siempre la última.'],
        ['• Al primer ingreso, cada persona debe definir su propia clave (mínimo 8 caracteres).'],
        ['• Quien ya definió su clave aparece como «Activa – clave propia» y no se modifica,'],
        ['  salvo que usted la restablezca desde la plataforma (botón 🔑).'],
        ['• Entregue cada clave solo a su titular y elimine esta planilla después de usarla.'],
    ]:
        ws.append(fila)
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 60

    ws = wb.create_sheet('Usuarios')
    ws.append(['Perfil', 'Región', 'Nombre', 'Correo (usuario)', 'Estado', 'Clave temporal', 'Secciones',
               'Ver', 'Agregar', 'Editar', 'Borrar', 'Solo su región', 'Último ingreso'])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='013722')
        c.alignment = Alignment(vertical='center')
    for f in filas:
        ws.append(f)
    for col, ancho in zip('ABCDEFGHIJKLM', [24, 20, 34, 34, 24, 16, 60, 6, 9, 8, 8, 13, 17]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    plano = io.BytesIO()
    wb.save(plano)
    plano.seek(0)
    cifrado = io.BytesIO()
    OOXMLFile(plano).encrypt(clave_admin, cifrado)
    resp = HttpResponse(cifrado.getvalue(),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="Permisos_SNAP_usuarios_{_dt.now():%Y-%m-%d_%H%M}.xlsx"'
    resp['Cache-Control'] = 'no-store'
    return resp


# ── Portal del investigador: accesos y registro geográfico ────────────────────

@api.post('/accesos')
@requiere(superusuario=True)
def accesos_generar(request):
    """Genera (o renueva) el acceso temporal del solicitante de un trámite.
    La clave solo se entrega en esta respuesta."""
    from .portal import generar_acceso
    try:
        tid = int(_cuerpo(request).get('tramite_id'))
        data = api_get(f'/tramites/{tid}')
        t = normalize_tramite(data.get('tramite', data))
    except (TypeError, ValueError):
        return api.create_response(request, {'ok': False, 'error': 'Número de trámite no válido.'}, status=400)
    except Exception:
        log.exception('No se pudo obtener el trámite %s para generar acceso', _cuerpo(request).get('tramite_id'))
        return api.create_response(request, {'ok': False, 'error': 'No se pudo consultar el trámite en Cero Filas.'}, status=502)
    if not t.get('id'):
        return api.create_response(request, {'ok': False, 'error': 'Trámite no encontrado.'}, status=404)
    acc, clave = generar_acceso(t, request.user)
    enlace = request.build_absolute_uri('/portal/ingresar')
    return {'ok': True, 'acceso': acc.como_dict(), 'clave': clave, 'enlace': enlace}


@api.get('/accesos')
@requiere(superusuario=True)
def accesos_listar(request):
    from .models import AccesoInvestigador, RegistroGeografico
    from django.db.models import Count
    n = dict(RegistroGeografico.objects.values_list('tramite_id').annotate(c=Count('id')))
    accs = AccesoInvestigador.objects.select_related('user').order_by('-creado_en')
    return {'ok': True, 'data': [{**a.como_dict(), 'registros': n.get(a.tramite_id, 0)} for a in accs]}


@api.post('/accesos/{acceso_id}/revocar')
@requiere(superusuario=True)
def accesos_revocar(request, acceso_id: int):
    from .models import AccesoInvestigador
    from .portal import revocar_acceso
    acc = AccesoInvestigador.objects.filter(pk=acceso_id).first()
    if not acc:
        return api.create_response(request, {'ok': False, 'error': 'Acceso no encontrado.'}, status=404)
    revocar_acceso(acc)
    return {'ok': True, 'acceso': acc.como_dict()}


def _registros_visibles(request, tramite_id):
    """Registros del trámite si el perfil puede ver su región (None si no)."""
    from .models import AccesoInvestigador, RegistroGeografico
    reg = _region_forzada(request)
    if reg:
        acc = AccesoInvestigador.objects.filter(tramite_id=tramite_id).first()
        if not acc or reg not in (acc.regiones or []):
            return None
    return RegistroGeografico.objects.filter(tramite_id=tramite_id)


@api.get('/tramites/{tramite_id}/registro-geografico')
def registro_geografico(request, tramite_id: int):
    from .models import AccesoInvestigador
    regs = _registros_visibles(request, tramite_id)
    if regs is None:
        return api.create_response(request, {'ok': False, 'error': 'Trámite fuera de su región.'}, status=403)
    acc = AccesoInvestigador.objects.filter(tramite_id=tramite_id).first()
    feats = [r.como_feature() for r in regs]
    if not privacidad.ve_sensible(request.user):
        feats = [privacidad.feature(f) for f in feats]
    return {'ok': True, 'type': 'FeatureCollection', 'features': feats,
            'acceso': acc.como_dict() if acc and request.user.is_superuser else None}


@api.get('/tramites/{tramite_id}/registro-geografico.geojson')
def registro_geografico_descarga(request, tramite_id: int):
    from .portal import geojson_descarga
    regs = _registros_visibles(request, tramite_id)
    if regs is None:
        return api.create_response(request, {'ok': False, 'error': 'Trámite fuera de su región.'}, status=403)
    return geojson_descarga(tramite_id, regs, proteger=not privacidad.ve_sensible(request.user))


# ── Vista territorial: país → región → área protegida ─────────────────────────

def _superposiciones(permisos, registros):
    """Pares de permisos cuyos registros geográficos se tocan (parcelas o transectos que se
    cruzan, o puntos a menos de ~100 m): señal de posible interferencia entre estudios."""
    from shapely.geometry import shape
    geoms = [(r.tramite_id, shape(r.geometria).buffer(0.001 if r.tipo == 'punto' else 0)) for r in registros]
    pares = set()
    for i, (a, ga) in enumerate(geoms):
        for b, gb in geoms[i + 1:]:
            if a != b and (min(a, b), max(a, b)) not in pares and ga.intersects(gb):
                pares.add((min(a, b), max(a, b)))
    return [{'a': a, 'b': b} for a, b in sorted(pares)]


def _proteger_areas(request, areas):
    """Sin rol Editor: coordenadas aproximadas (~1 km) y textos sin datos personales."""
    if privacidad.ve_sensible(request.user):
        return areas
    import copy
    areas = copy.deepcopy(areas)
    for a in areas:
        for p in a.get('permisos', []):
            p['puntos'] = privacidad.puntos(p.get('puntos') or [])
            p['titulo'] = privacidad.redactar(p.get('titulo') or '')
            p['registros'] = [privacidad.feature(f) for f in p.get('registros') or []]
    return areas


@api.get('/territorio')
def territorio(request, tipo: Optional[str] = None, region: Optional[str] = None):
    """Concentración de permisos activos. Sin región: todas las áreas del país (centro y
    superficie). Con región: sus áreas con límite, permisos, coordenadas declaradas, el
    registro geográfico del portal del investigador y las superposiciones entre permisos."""
    from django.core.cache import cache
    from .models import RegistroGeografico
    from .presion import presion_nacional, presion_region
    if tipo not in (None, 'inv', 'film'):
        return api.create_response(request, {'ok': False, 'error': 'tipo debe ser inv o film'}, status=400)
    forzada = _region_forzada(request)
    if forzada and region != forzada:
        return api.create_response(request, {'ok': False, 'error': 'Solo puede consultar su región.'}, status=403)
    try:
        if not region:
            clave = f'territorio:nac:v3:{tipo}:{datetime.now():%Y%m%d%H}'
            datos = cache.get(clave)
            if datos is None:
                datos = presion_nacional(tipo)
                cache.set(clave, datos, 3600)
            return {'ok': True, 'escala': 'pais', 'tipo': tipo, 'areas': datos}
        areas = presion_region(region, tipo)
    except Exception:
        log.exception('Error al calcular la vista territorial (%s, %s)', region, tipo)
        return api.create_response(request, {'ok': False, 'error': 'No se pudo calcular la vista territorial.'}, status=500)
    ids = {p['id'] for a in areas for p in a['permisos']}
    regs = {}
    for r in RegistroGeografico.objects.filter(tramite_id__in=ids):
        regs.setdefault(r.tramite_id, []).append(r)
    for a in areas:
        propios = [r for p in a['permisos'] for r in regs.get(p['id'], [])]
        for p in a['permisos']:
            p['registros'] = [r.como_feature() for r in regs.get(p['id'], [])]
        a['superposiciones'] = _superposiciones(a['permisos'], propios)
    return {'ok': True, 'escala': 'region', 'tipo': tipo, 'region': region, 'areas': _proteger_areas(request, areas)}


# ── Búsqueda: índice de todos los años ────────────────────────────────────────
BUSQUEDA_DESDE = 2019
_CAMPOS_TEXTO = ('objetivos_del_proyecto', 'descripcion', 'descripcion_proyecto', 'segmentos_proyecto',
                 'objetivo_general', 'objetivos_especificos', 'actividades', 'actividades_a_realizar',
                 'palabras_clave', 'especie', 'especies', 'nombre_cientifico', 'nombre_comun',
                 'tipo_investigacion', 'linea_investigacion')


def _anios_busqueda():
    """Trámites de cada año desde BUSQUEDA_DESDE, solo de los años ya en caché. Los que
    faltan se descargan en segundo plano (la primera vez tardan) y se informan."""
    from . import tramites_cache as tc
    hoy = datetime.now(tc.TZ).year
    datos, pendientes = [], []
    for a in range(BUSQUEDA_DESDE, hoy + 1):
        entrada = tc.cache.get(tc._clave(a, a))
        if entrada is None:
            tc._refrescar_en_segundo_plano(a, a)
            pendientes.append(a)
            continue
        if time.time() - entrada['ts'] > tc.FRESCO_SEG:
            tc._refrescar_en_segundo_plano(a, a)
        datos.extend(entrada['data'])
    return datos, pendientes


def _registro_busqueda(t, con_registro):
    from .core import extract_titulo
    d = t.get('datos_raw') or {}
    ea = t.get('etapa_actual') or {}
    texto = ' '.join(str(d.get(k)) for k in _CAMPOS_TEXTO if d.get(k))
    return {
        'id': t['id'], 'proceso_id': t.get('proceso_id'), 'estado': t.get('estado'),
        'recepcion_pendiente': t.get('recepcion_pendiente'), 'suspendido': t.get('suspendido'),
        'espera_solicitante': t.get('espera_solicitante'),
        'fecha_inicio': t.get('fecha_inicio'), 'fecha_resolucion': t.get('fecha_resolucion'),
        'fecha_termino': t.get('fecha_termino'), 'fecha_modificacion': t.get('fecha_modificacion'),
        'fecha_limite_proceso': t.get('fecha_limite_proceso'),
        'nombre_solicitante': t.get('nombre_solicitante') or '', 'email_solicitante': t.get('email_solicitante') or '',
        'titulo_investigacion': (extract_titulo(d) or '')[:400],
        'texto_clasificacion': texto[:800],
        'areas_protegidas': t.get('areas_protegidas') or [], 'regiones_list': t.get('regiones_list') or [],
        # sin nombres de funcionarios: solo la etapa, su plazo y si está en CONAF
        'etapa_actual': {'tarea_nombre': ea.get('tarea_nombre'), 'fecha_vencimiento': ea.get('fecha_vencimiento'),
                         'fecha_inicio': ea.get('fecha_inicio'), 'email': ea.get('email') or ''} if ea else None,
        'etapas_completadas': t.get('etapas_completadas'),
        'porcentaje_avance': t.get('porcentaje_avance'),
        'archivo_autorizacion': bool(t.get('archivo_autorizacion')), 'etapa_autorizacion_id': t.get('etapa_autorizacion_id'),
        'tiene_registro': t['id'] in con_registro,
    }


@api.get('/buscar/indice')
@requiere(seccion='busqueda')
def buscar_indice(request):
    """Registros livianos de todos los años para buscar y filtrar en el navegador."""
    import gzip
    from django.core.cache import cache
    from .models import RegistroGeografico
    reg = _region_forzada(request) or 'todas'
    protegido = not privacidad.ve_sensible(request.user)
    clave = f'buscar:indice:v2:{reg}:{"protegido" if protegido else "completo"}'
    cuerpo = cache.get(clave)
    if cuerpo is None:
        datos, pendientes = _anios_busqueda()
        datos = _filtrar_region(request, [t for t in datos if es_tramite_visible(t)])
        if protegido:
            datos = [privacidad.tramite(t) for t in datos]
        con_registro = set(RegistroGeografico.objects.values_list('tramite_id', flat=True).distinct())
        vistos, registros = set(), []
        for t in datos:
            if t['id'] not in vistos:
                vistos.add(t['id'])
                registros.append(_registro_busqueda(t, con_registro))
        cuerpo = gzip.compress(_json.dumps({'ok': True, 'data': registros, 'pendientes': pendientes,
                                            'desde': BUSQUEDA_DESDE}, ensure_ascii=False).encode())
        if not pendientes:        # incompleto: no se guarda, para reintentar cuando terminen las descargas
            cache.set(clave, cuerpo, 600)
    if 'gzip' in request.headers.get('Accept-Encoding', ''):
        resp = HttpResponse(cuerpo, content_type='application/json')
        resp['Content-Encoding'] = 'gzip'
    else:
        resp = HttpResponse(gzip.decompress(cuerpo), content_type='application/json')
    resp['Vary'] = 'Accept-Encoding, Cookie'
    resp['Cache-Control'] = 'private, max-age=300'
    return resp


@api.post('/buscar/excel')
@requiere(seccion='exportar')
def buscar_excel(request):
    """Excel con los trámites de una búsqueda (lista de IDs), respetando la región del perfil."""
    try:
        ids = {int(i) for i in (_cuerpo(request).get('ids') or [])[:5000]}
    except (TypeError, ValueError):
        return api.create_response(request, {'ok': False, 'error': 'Lista de trámites no válida.'}, status=400)
    datos, _ = _anios_busqueda()
    vistos, tramites = set(), []
    for t in _filtrar_region(request, datos):
        if t['id'] in ids and t['id'] not in vistos and es_tramite_visible(t):
            vistos.add(t['id'])
            tramites.append(t)
    buf = build_tramites_excel(privacidad.tramites(request, tramites), {})
    resp = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="busqueda_{datetime.now():%Y%m%d_%H%M}.xlsx"'
    return resp


# ── Adjuntos servidos por la plataforma (sin cuenta en Cero Filas) ────────────

def _tramite_con_acceso(request, tramite_id):
    """Trámite si el perfil puede verlo; si no, la respuesta de error."""
    from .adjuntos import tramite_vivo
    try:
        t = tramite_vivo(tramite_id)
    except Exception:
        log.exception('No se pudo consultar el trámite %s', tramite_id)
        return None, api.create_response(request, {'ok': False, 'error': 'No se pudo consultar el trámite en Cero Filas.'}, status=502)
    if not t.get('id') or not es_tramite_visible(t):
        return None, api.create_response(request, {'ok': False, 'error': 'Trámite no encontrado.'}, status=404)
    if not _filtrar_region(request, [t]):
        return None, api.create_response(request, {'ok': False, 'error': 'Trámite fuera de su región.'}, status=403)
    return t, None


@api.get('/tramites/{tramite_id}/adjuntos.zip')
def adjuntos_zip(request, tramite_id: int):
    from .adjuntos import adjuntos_de, zip_tramite
    t, err = _tramite_con_acceso(request, tramite_id)
    if err:
        return err
    if not adjuntos_de(t):
        return api.create_response(request, {'ok': False, 'error': 'El trámite no tiene adjuntos.'}, status=404)
    resp = HttpResponse(zip_tramite(t, request.user, proteger=not privacidad.ve_sensible(request.user)), content_type='application/zip')
    resp['Content-Disposition'] = f'attachment; filename="adjuntos_{int(tramite_id)}.zip"'
    return resp


@api.get('/tramites/{tramite_id}/adjuntos/{clave}')
def adjunto(request, tramite_id: int, clave: str):
    from .adjuntos import AdjuntoError, es_adjunto, respuesta
    t, err = _tramite_con_acceso(request, tramite_id)
    if err:
        return err
    archivo = (t.get('datos_raw') or {}).get(clave)
    if not es_adjunto(archivo):
        return api.create_response(request, {'ok': False, 'error': 'Adjunto no encontrado.'}, status=404)
    try:
        return respuesta(tramite_id, clave, archivo, request.user, proteger=not privacidad.ve_sensible(request.user))
    except AdjuntoError as e:
        if e.status == 403:      # se abre en una pestaña o en el visor: página legible con el aviso legal
            from django.shortcuts import render
            return render(request, 'dato_protegido.html', {'mensaje': str(e), 'tramite_id': tramite_id}, status=403)
        return api.create_response(request, {'ok': False, 'error': str(e)}, status=e.status)


@api.post('/buscar/resoluciones.zip')
@requiere(seccion='exportar')
def buscar_resoluciones_zip(request):
    """ZIP con las resoluciones de una lista de trámites (máx. 150), respetando la región del perfil."""
    from .adjuntos import MAX_ZIP_TRAMITES, zip_resoluciones
    try:
        ids = [int(i) for i in (_cuerpo(request).get('ids') or [])]
    except (TypeError, ValueError):
        return api.create_response(request, {'ok': False, 'error': 'Lista de trámites no válida.'}, status=400)
    pedidos = set(ids[:MAX_ZIP_TRAMITES])
    datos, _ = _anios_busqueda()
    vistos, tramites = set(), []
    for t in _filtrar_region(request, datos):
        if t['id'] in pedidos and t['id'] not in vistos and es_tramite_visible(t):
            vistos.add(t['id'])
            tramites.append(t)
    resp = HttpResponse(zip_resoluciones(privacidad.tramites(request, tramites), request.user,
                                         proteger=not privacidad.ve_sensible(request.user)), content_type='application/zip')
    resp['Content-Disposition'] = f'attachment; filename="resoluciones_{datetime.now():%Y%m%d_%H%M}.zip"'
    return resp
